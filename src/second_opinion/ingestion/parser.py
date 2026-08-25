from __future__ import annotations

import hashlib
import uuid
from pathlib import PurePosixPath

from ..domain.documents import Document

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".docx", ".pdf"}


class ParseError(ValueError):
    pass


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u2011", "-")
    return text


def _new_document(filename: str, content_type: str, text: str, page_offsets: list[int]) -> Document:
    normalized = _normalize_text(text)
    if not normalized.strip():
        raise ParseError("документ не содержит текста")
    return Document(
        document_id=uuid.uuid4().hex,
        filename=filename,
        content_type=content_type,
        text=normalized,
        sha256=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        page_offsets=page_offsets,
    )


def parse_txt(filename: str, content: bytes) -> Document:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("cp1251", errors="replace")
    return _new_document(filename, "text/plain", text, [])


def parse_markdown(filename: str, content: bytes) -> Document:
    doc = parse_txt(filename, content)
    return doc.model_copy(update={"content_type": "text/markdown"})


def parse_docx(filename: str, content: bytes) -> Document:
    try:
        import io

        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise ParseError("python-docx не установлен") from exc
    try:
        docx = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise ParseError(f"не удалось прочитать DOCX: {exc}") from exc
    text = "\n".join(p.text for p in docx.paragraphs)
    return _new_document(filename, "application/docx", text, [])


def parse_pdf(filename: str, content: bytes) -> Document:
    try:
        import io

        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParseError("pypdf не установлен") from exc
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ParseError(f"не удалось прочитать PDF: {exc}") from exc
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    if not any(p.strip() for p in pages):
        raise ParseError("PDF не содержит текстового слоя (скан без OCR)")
    page_offsets: list[int] = []
    cursor = 0
    for page_text in pages:
        page_offsets.append(cursor)
        cursor += len(_normalize_text(page_text)) + 1
    return _new_document(filename, "application/pdf", "\n".join(pages), page_offsets)


PARSERS = {
    ".txt": parse_txt,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".docx": parse_docx,
    ".pdf": parse_pdf,
}


def parse_document(filename: str, content: bytes, max_bytes: int) -> Document:
    """Единая точка входа парсинга: выбор парсера по расширению."""
    if len(content) == 0:
        raise ParseError("файл пуст")
    if len(content) > max_bytes:
        raise ParseError(
            f"файл больше лимита {max_bytes} байт (получено {len(content)})"
        )
    suffix = PurePosixPath(filename).suffix.lower()
    parser = PARSERS.get(suffix)
    if parser is None:
        raise ParseError(
            f"неподдерживаемый формат '{suffix or 'без расширения'}'; "
            f"поддерживаются: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return parser(filename, content)
