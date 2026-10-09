"""Заключение в формате .docx (ФТ-34 из docs/TZ.md).

Файл можно приложить к материалам и переслать, лист из принтера нельзя.

Документ собирается из листа заключения (``report_sheet``), то есть из готового
отчёта, и ничего к нему не добавляет: те же разделы и таблицы, что на экране и
в печатной форме. Выгрузка не должна стать вторым, неконтролируемым
представлением результата.

Зависимость `python-docx` уже есть в проекте: ею читаются загружаемые
документы. Новых зависимостей выгрузка не требует.
"""

from __future__ import annotations

import io

from docx import Document as new_document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..domain.analysis import AnalysisReport
from ..domain.documents import Document
from .report_sheet import Preparer, Sheet, Table, build_sheet

USABLE_WIDTH_CM = 26.7
TABLE_FONT_PT = 9


def _plain_heading_style(doc: DocxDocument, name: str) -> None:
    style = doc.styles[name]
    if style.font.color is not None:
        style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    properties = style.element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _write_cell(cell, text: str, bold: bool = False) -> None:
    lines = text.split("\n")
    first = cell.paragraphs[0]
    for index, line in enumerate(lines):
        paragraph = first if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.size = Pt(TABLE_FONT_PT)
        run.bold = bold


def _add_table(doc: DocxDocument, table: Table) -> None:
    columns = len(table.headers)
    widths = table.widths or [1.0] * columns
    total = sum(widths)
    grid = doc.add_table(rows=1, cols=columns)
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for index, header in enumerate(table.headers):
        cell = grid.rows[0].cells[index]
        _write_cell(cell, header, bold=True)
        _shade(cell, "E7E6E6")
    for row in table.rows:
        cells = grid.add_row().cells
        for index, value in enumerate(row):
            _write_cell(cells[index], value, bold=table.key_value and index == 0)
    for row in grid.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(USABLE_WIDTH_CM * widths[index] / total)
    header_row = grid.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_row.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_docx(sheet: Sheet) -> bytes:
    doc = new_document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(1.5))

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for name in ("Title", "Heading 1"):
        _plain_heading_style(doc, name)

    title = doc.add_heading(sheet.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in sheet.blocks:
        if block.kind == "heading":
            doc.add_heading(block.text, level=1)
        elif block.kind == "table" and block.table is not None:
            _add_table(doc, block.table)
        else:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block.text)
            run.bold = block.kind == "lead"
            run.italic = block.kind == "note"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_docx(
    report: AnalysisReport,
    document: Document | None = None,
    mask_personal_data: bool = False,
    prepared_by: Preparer | None = None,
) -> bytes:
    """Собрать заключение в .docx и вернуть его байтами.

    ``mask_personal_data`` скрывает персональные данные в цитатах и в изложении
    (Д-06). Оригинал документа при этом не меняется: обезличивается то, что
    уходит из системы, а не то, что хранится.
    """
    return render_docx(build_sheet(report, document, mask_personal_data, prepared_by))
