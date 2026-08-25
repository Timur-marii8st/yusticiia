"""Однократная генерация бинарных фикстур для тестов парсера.

Запуск:  python scripts/generate_test_fixtures.py
Результат коммитится в tests/fixtures/.
"""

from __future__ import annotations

import io
from pathlib import Path

FIXTURES = Path(__file__).resolve().parents[1] / "tests" / "fixtures"

PDF_TEXT = "Prigovor: ch. 2 st. 158 UK RF, lishenie svobody na srok 2 goda."


def build_pdf(text: str) -> bytes:
    stream = f"BT /F1 14 Tf 50 780 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode("ascii"))
        out.write(obj)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(b"xref\n")
    out.write(f"0 {len(objects) + 1}\n".encode("ascii"))
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    out.write(b"trailer\n")
    out.write(f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode("ascii"))
    out.write(b"startxref\n")
    out.write(f"{xref_pos}\n".encode("ascii"))
    out.write(b"%%EOF")
    return out.getvalue()


def build_docx() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("ПРИГОВОР (синтетический образец)")
    doc.add_paragraph("Действия подсудимого квалифицированы по ч. 2 ст. 158 УК РФ.")
    doc.add_paragraph("Назначить наказание в виде лишения свободы на срок 2 года.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    (FIXTURES / "sample_text.pdf").write_bytes(build_pdf(PDF_TEXT))
    (FIXTURES / "sample_text.docx").write_bytes(build_docx())
    print(f"фикстуры созданы в {FIXTURES}")


if __name__ == "__main__":
    main()
