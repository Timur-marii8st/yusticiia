"""Сборка концепции и описания архитектуры в DOCX и PDF из Markdown.

Источник текста — ``docs/CONCEPT.md`` (``--doc concept``, по умолчанию) или
``docs/ARCHITECTURE.md`` (``--doc architecture``); рисунки — ``docs/concept/*.png``
(их рисует ``scripts/concept_figures.py``). Документы для отправки на конкурс
собираются из них, поэтому текст в репозитории и в отправленном файле не
расходятся.

Пример:
  python scripts/concept_figures.py
  python scripts/build_concept.py --out build/concept
  python scripts/build_concept.py --doc architecture --out build/concept
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TITLE = "«Второе мнение»"
CONTEST = "Всероссийский конкурс студенческих проектов\n«Алгоритм правосудия. Второе мнение»"
FONTS = Path("/System/Library/Fonts/Supplemental")


@dataclass(frozen=True)
class Spec:
    """Что собирается: источник, имя файлов, подзаголовок и надпись титульного листа."""

    source: Path
    basename: str
    subtitle: str
    cover: str
    pdf_title: str


SPECS = {
    "concept": Spec(
        ROOT / "docs" / "CONCEPT.md",
        "Концепция_Второе_мнение",
        "Система проверки проекта приговора в части назначения наказания",
        "Концепция проекта (конкурсная работа отборочного этапа)\n" + CONTEST,
        f"Концепция проекта {TITLE}",
    ),
    "architecture": Spec(
        ROOT / "docs" / "ARCHITECTURE.md",
        "Архитектура_Второе_мнение",
        "Архитектура решения",
        "Описание архитектуры решения к концепции проекта\n" + CONTEST,
        f"Архитектура решения {TITLE}",
    ),
}

_MONTHS = ("января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
           "сентября", "октября", "ноября", "декабря")


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    path: Path | None = None


def parse(markdown: str, base: Path) -> list[Block]:
    """Абзацы, заголовки, списки, таблицы, рисунки и блоки кода; пути рисунков — от ``base``."""
    blocks: list[Block] = []
    lines = markdown.splitlines()
    i = 0
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("p", " ".join(s.strip() for s in paragraph)))
            paragraph.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        if stripped.startswith("```"):
            flush()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            blocks.append(Block("code", "\n".join(code)))
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if heading:
            flush()
            blocks.append(Block("h", heading.group(2), level=len(heading.group(1))))
        elif image:
            flush()
            blocks.append(Block("img", image.group(1), path=base / image.group(2)))
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        elif re.match(r"^(-|\d+\.)\s+", stripped):
            flush()
            kind = "ul" if stripped.startswith("-") else "ol"
            number = re.match(r"^(\d+)\.", stripped)
            text = re.sub(r"^(-|\d+\.)\s+", "", stripped)
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                text += " " + lines[i].strip()
                i += 1
            blocks.append(Block(kind, text, level=int(number.group(1)) if number else 0))
            continue
        else:
            paragraph.append(stripped)
        i += 1
    flush()
    return blocks


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def segments(text: str) -> list[tuple[str, str]]:
    """Разбить строку на куски: обычный текст, полужирный, код."""
    result: list[tuple[str, str]] = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            result.append(("bold", part[2:-2]))
        elif part.startswith("`"):
            result.append(("code", part[1:-1]))
        elif part.startswith("["):
            result.append(("plain", re.sub(r"^\[([^\]]+)\]\([^)]+\)$", r"\1", part)))
        else:
            result.append(("plain", part))
    return result


def plain(text: str) -> str:
    """Текст без разметки: для оглавления и заголовков DOCX."""
    return "".join(chunk for _, chunk in segments(text))


def _ru_date(day: date) -> str:
    return f"{day.day} {_MONTHS[day.month - 1]} {day.year} г."


def _body_without_title(blocks: list[Block]) -> list[Block]:
    """Заголовок и подзаголовок из Markdown уходят на титульный лист."""
    body = list(blocks)
    if body and body[0].kind == "h" and body[0].level == 1:
        body = body[1:]
    if body and body[0].kind == "p" and body[0].text.startswith("**"):
        body = body[1:]
    return body


def _column_widths(rows: list[list[str]], total: float) -> list[float]:
    """Ширины столбцов в см: не уже самого длинного слова, остаток — по объёму текста."""
    columns = max(len(r) for r in rows)
    cells = [[re.sub(r"[*`]", "", r[c]) if c < len(r) else "" for r in rows] for c in range(columns)]
    minimum = [max((len(word) for text in col for word in text.split()), default=4) * 0.2 + 0.4 for col in cells]
    weights = [min(max(max(len(t) for t in col), 8), 90) ** 0.8 for col in cells]
    if sum(minimum) >= total:
        return [total * m / sum(minimum) for m in minimum]
    spare = total - sum(minimum)
    return [m + spare * w / sum(weights) for m, w in zip(minimum, weights, strict=True)]


def build_docx(blocks: list[Block], target: Path, built: date, spec: Spec) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    usable = 17.0

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), "Times New Roman")
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    def runs(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
        for kind, chunk in segments(text):
            run = paragraph.add_run(chunk)
            run.bold = bold or kind == "bold"
            if kind == "code":
                run.font.name = "Courier New"
                run.font.size = Pt((size or 12) - 1.5)
            elif size:
                run.font.size = Pt(size)

    for _ in range(6):
        doc.add_paragraph()
    for text, size, bold in ((TITLE, 26, True), (spec.subtitle, 15, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
    for _ in range(3):
        doc.add_paragraph()
    for line in spec.cover.split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(13)
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Редакция от {_ru_date(built)}").font.size = Pt(12)

    body = _body_without_title(blocks)
    heading = doc.add_paragraph("Содержание", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.page_break_before = True
    for block in body:
        if block.kind == "h" and block.level == 2:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(plain(block.text))
        elif block.kind == "h" and block.level == 3:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(0)
            p.add_run(plain(block.text)).font.size = Pt(11)
    first = True
    for block in body:
        if block.kind == "h":
            added = doc.add_heading(plain(block.text), level=max(1, block.level - 1))
            if first:
                added.paragraph_format.page_break_before = True
                first = False
        elif block.kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, block.text)
            if first:
                p.paragraph_format.page_break_before = True
                first = False
        elif block.kind == "ul":
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, block.text)
        elif block.kind == "ol":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(f"{block.level}. ")
            runs(p, block.text)
        elif block.kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(6)
            code_lines = block.text.split("\n")
            for index, line in enumerate(code_lines):
                run = p.add_run(line)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                if index < len(code_lines) - 1:
                    run.add_break()
        elif block.kind == "img" and block.path is not None:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
            p.add_run().add_picture(str(block.path), width=Cm(16.5))
        elif block.kind == "table":
            columns = max(len(r) for r in block.rows)
            widths = _column_widths(block.rows, usable)
            table = doc.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for c in range(columns):
                    cell = cells[c]
                    cell.width = Cm(widths[c])
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    runs(paragraph, row[c] if c < len(row) else "", size=10, bold=index == 0)
                    if index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:val"), "clear")
                        shading.set(qn("w:color"), "auto")
                        shading.set(qn("w:fill"), "E7E6E6")
                        cell._tc.get_or_add_tcPr().append(shading)
            header = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            header.append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field_begin, instr, field_end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    field_end.set(qn("w:fldCharType"), "end")
    run = footer.add_run()
    run.font.size = Pt(10)
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_end)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target


def build_pdf(blocks: list[Block], target: Path, built: date, spec: Spec) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    faces = {
        "TNR": "Times New Roman.ttf",
        "TNR-Bold": "Times New Roman Bold.ttf",
        "TNR-Italic": "Times New Roman Italic.ttf",
        "TNR-BoldItalic": "Times New Roman Bold Italic.ttf",
        "Mono": "Courier New.ttf",
    }
    for name, file in faces.items():
        pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-Bold", italic="TNR-Italic", boldItalic="TNR-BoldItalic")

    base = ParagraphStyle("base", fontName="TNR", fontSize=11.5, leading=14.5, alignment=TA_JUSTIFY, spaceAfter=4)
    cell = ParagraphStyle("cell", parent=base, fontSize=9.5, leading=11.5, alignment=0, spaceAfter=0)
    head = {
        2: ParagraphStyle("h2", parent=base, fontName="TNR-Bold", fontSize=15, leading=18, spaceBefore=12, spaceAfter=6, alignment=0, keepWithNext=1),
        3: ParagraphStyle("h3", parent=base, fontName="TNR-Bold", fontSize=12.5, leading=15, spaceBefore=8, spaceAfter=4, alignment=0, keepWithNext=1),
    }
    center = ParagraphStyle("center", parent=base, alignment=TA_CENTER)

    def markup(text: str) -> str:
        out = []
        for kind, chunk in segments(text):
            chunk = escape(chunk)
            if kind == "bold":
                out.append(f"<b>{chunk}</b>")
            elif kind == "code":
                out.append(f'<font name="Mono" size="9.5">{chunk}</font>')
            else:
                out.append(chunk)
        return "".join(out)

    story: list = [Spacer(1, 5 * cm)]
    story.append(Paragraph(escape(TITLE), ParagraphStyle("t", parent=center, fontName="TNR-Bold", fontSize=26, leading=32)))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph(escape(spec.subtitle), ParagraphStyle("st", parent=center, fontSize=15, leading=19)))
    story.append(Spacer(1, 2.5 * cm))
    for line in spec.cover.split("\n"):
        story.append(Paragraph(escape(line), ParagraphStyle("c", parent=center, fontSize=13, leading=17)))
    story.append(Spacer(1, 7 * cm))
    story.append(Paragraph(f"Редакция от {_ru_date(built)}", center))
    story.append(PageBreak())

    body = _body_without_title(blocks)
    story.append(Paragraph("Содержание", head[2]))
    for block in body:
        if block.kind == "h" and block.level in (2, 3):
            indent = 0 if block.level == 2 else 0.8 * cm
            story.append(Paragraph(escape(plain(block.text)), ParagraphStyle("toc", parent=base, leftIndent=indent, spaceAfter=1, fontSize=11 if block.level == 2 else 10.5, leading=13)))
    story.append(PageBreak())

    usable = A4[0] - 4.0 * cm
    for block in body:
        if block.kind == "h":
            story.append(Paragraph(markup(block.text), head[min(3, max(2, block.level))]))
        elif block.kind == "p":
            story.append(Paragraph(markup(block.text), base))
        elif block.kind == "ul":
            story.append(Paragraph(markup(block.text), ParagraphStyle("li", parent=base, leftIndent=0.9 * cm, bulletIndent=0.3 * cm), bulletText="•"))
        elif block.kind == "ol":
            story.append(Paragraph(markup(block.text), ParagraphStyle("ol", parent=base, leftIndent=0.9 * cm, bulletIndent=0.2 * cm), bulletText=f"{block.level}."))
        elif block.kind == "code":
            story.append(Preformatted(block.text, ParagraphStyle(
                "code", fontName="Mono", fontSize=8.5, leading=10.5, leftIndent=0.3 * cm, spaceBefore=4, spaceAfter=8,
            )))
        elif block.kind == "img" and block.path is not None:
            from PIL import Image as PILImage

            with PILImage.open(block.path) as picture:
                width, height = picture.size
            draw_width = usable
            draw_height = draw_width * height / width
            limit = A4[1] - 6 * cm
            if draw_height > limit:
                draw_height, draw_width = limit, limit * width / height
            story.append(KeepTogether([Spacer(1, 0.2 * cm), Image(str(block.path), width=draw_width, height=draw_height), Spacer(1, 0.3 * cm)]))
        elif block.kind == "table":
            widths = [w * cm for w in _column_widths(block.rows, usable / cm)]
            columns = len(widths)
            data = [
                [Paragraph(("<b>%s</b>" if r == 0 else "%s") % markup(row[c] if c < len(row) else ""), cell) for c in range(columns)]
                for r, row in enumerate(block.rows)
            ]
            table = Table(data, colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#808080")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7E6E6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(Spacer(1, 0.15 * cm))
            story.append(table)
            story.append(Spacer(1, 0.3 * cm))

    def number(canvas, document) -> None:
        if document.page > 1:
            canvas.setFont("TNR", 10)
            canvas.drawCentredString(A4[0] / 2, 1.2 * cm, str(document.page))

    target.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(target), pagesize=A4, leftMargin=2.5 * cm, rightMargin=1.5 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=spec.pdf_title, author="Команда проекта «Второе мнение»", lang="ru-RU",
    )
    pdf.build(story, onFirstPage=number, onLaterPages=number)
    return target


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, help="каталог для DOCX и PDF")
    parser.add_argument("--doc", choices=sorted(SPECS), default="concept", help="что собирать (по умолчанию concept)")
    parser.add_argument("--date", default=None, help="дата редакции ГГГГ-ММ-ДД (по умолчанию сегодня)")
    args = parser.parse_args(argv)
    built = date.fromisoformat(args.date) if args.date else date.today()
    spec = SPECS[args.doc]
    blocks = parse(spec.source.read_text(encoding="utf-8"), spec.source.parent)
    missing = [b.path for b in blocks if b.kind == "img" and b.path is not None and not b.path.exists()]
    if missing:
        print(f"ОШИБКА: нет рисунков {missing}; сначала python scripts/concept_figures.py", file=sys.stderr)
        return 2
    out = Path(args.out)
    for path in (
        build_docx(blocks, out / f"{spec.basename}.docx", built, spec),
        build_pdf(blocks, out / f"{spec.basename}.pdf", built, spec),
    ):
        print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
